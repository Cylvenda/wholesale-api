from datetime import date, datetime, time, timedelta
from decimal import Decimal, ROUND_HALF_UP
from html import escape
from io import BytesIO
from uuid import UUID

from django.db.models import Prefetch
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from rest_framework.exceptions import ValidationError

from .models import BusinessDetails
from ..customers.models import Customer
from ..products.models import Product
from ..purchases.models import Purchase, PurchaseItem
from ..sales.models import Sale, SaleItem
from ..suppliers.models import Supplier


MONEY_FORMAT = "#,##0.00"
DATE_FORMAT = "yyyy-mm-dd"
HEADER_FILL = PatternFill("solid", fgColor="1D4ED8")
HEADER_FONT = Font(color="FFFFFF", bold=True)
THIN_BORDER = Border(bottom=Side(style="thin", color="D1D5DB"))
CENT = Decimal("0.01")


def parse_report_date(value, field_name):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    try:
        return date.fromisoformat(str(value))
    except (TypeError, ValueError) as exc:
        raise ValidationError({field_name: "Use a valid date in YYYY-MM-DD format."}) from exc


def validate_date_range(from_date, to_date):
    if from_date and to_date and from_date > to_date:
        raise ValidationError({"from": "The from date must be on or before the to date."})


def validate_uuid(value, field_name, model):
    if value in (None, ""):
        return None

    try:
        parsed = value if isinstance(value, UUID) else UUID(str(value))
    except (TypeError, ValueError) as exc:
        raise ValidationError({field_name: "Enter a valid UUID."}) from exc

    if not model.objects.filter(uuid=parsed).exists():
        raise ValidationError({field_name: "The selected record does not exist."})

    return parsed


def timezone_bounds(from_date, to_date):
    validate_date_range(from_date, to_date)
    current_timezone = timezone.get_current_timezone()

    start = None
    if from_date:
        start = timezone.make_aware(
            datetime.combine(from_date, time.min),
            current_timezone,
        )

    end = None
    if to_date:
        end = timezone.make_aware(
            datetime.combine(to_date + timedelta(days=1), time.min),
            current_timezone,
        )

    return start, end


def style_worksheet(worksheet):
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_view.showGridLines = False

    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    worksheet.row_dimensions[1].height = 24

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for column_cells in worksheet.columns:
        values = [str(cell.value) if cell.value is not None else "" for cell in column_cells]
        width = min(max(len(value) for value in values) + 2, 48)
        worksheet.column_dimensions[get_column_letter(column_cells[0].column)].width = max(width, 12)


def save_workbook(workbook):
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def report_filename(prefix, from_date, to_date):
    start = from_date.isoformat() if from_date else "all"
    end = to_date.isoformat() if to_date else "present"
    return f"{prefix}-{start}-to-{end}.xlsx"


def excel_response(workbook, filename):
    response = HttpResponse(
        save_workbook(workbook).getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def allocate_discount(items, discount):
    discount = discount or Decimal("0.00")
    if discount <= 0 or not items:
        return [Decimal("0.00") for _ in items]

    def get_subtotal(item):
        return item.subtotal if hasattr(item, "subtotal") else item["subtotal"]

    subtotal = sum((get_subtotal(item) for item in items), Decimal("0.00"))
    if subtotal <= 0:
        return [Decimal("0.00") for _ in items]

    allocations = []
    allocated = Decimal("0.00")
    for item in items[:-1]:
        amount = (discount * get_subtotal(item) / subtotal).quantize(CENT, rounding=ROUND_HALF_UP)
        amount = min(amount, discount - allocated)
        allocations.append(amount)
        allocated += amount

    allocations.append(discount - allocated)
    return allocations


def build_purchases_excel(*, from_date=None, to_date=None, product_uuid=None, supplier_uuid=None):
    from_date = parse_report_date(from_date, "from")
    to_date = parse_report_date(to_date, "to")
    product_id = validate_uuid(product_uuid, "product", Product)
    supplier_id = validate_uuid(supplier_uuid, "supplier", Supplier)
    start, end = timezone_bounds(from_date, to_date)

    items = PurchaseItem.objects.select_related("product").order_by("purchase__purchase_date", "pk")
    purchases = (
        Purchase.objects
        .filter(status=Purchase.Status.COMPLETED)
        .select_related("supplier")
        .prefetch_related(Prefetch("items", queryset=items, to_attr="report_items"))
        .order_by("purchase_date", "pk")
    )

    if start:
        purchases = purchases.filter(purchase_date__gte=start)
    if end:
        purchases = purchases.filter(purchase_date__lt=end)
    if product_id:
        purchases = purchases.filter(items__product_id=product_id).distinct()
    if supplier_id:
        purchases = purchases.filter(supplier_id=supplier_id)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Purchased Items"
    worksheet.append([
        "Date",
        "Purchase No",
        "Supplier",
        "Product",
        "Quantity",
        "Unit Cost",
        "Total",
    ])

    for purchase in purchases:
        purchase_number = purchase.invoice_number or f"PUR-{purchase.uuid}"
        for item in purchase.report_items:
            worksheet.append([
                timezone.localtime(purchase.purchase_date).date(),
                purchase_number,
                purchase.supplier.name,
                item.product.name,
                item.quantity,
                item.unit_cost,
                item.subtotal,
            ])
            worksheet.cell(worksheet.max_row, 1).number_format = DATE_FORMAT
            worksheet.cell(worksheet.max_row, 5).number_format = "0"
            worksheet.cell(worksheet.max_row, 6).number_format = MONEY_FORMAT
            worksheet.cell(worksheet.max_row, 7).number_format = MONEY_FORMAT

    style_worksheet(worksheet)
    return excel_response(workbook, report_filename("purchased-items", from_date, to_date))


def build_sales_excel(*, from_date=None, to_date=None, product_uuid=None, customer_uuid=None):
    from_date = parse_report_date(from_date, "from")
    to_date = parse_report_date(to_date, "to")
    product_id = validate_uuid(product_uuid, "product", Product)
    customer_id = validate_uuid(customer_uuid, "customer", Customer)
    start, end = timezone_bounds(from_date, to_date)

    items = SaleItem.objects.select_related("product").order_by("sale__sale_date", "pk")
    sales = (
        Sale.objects
        .filter(status=Sale.Status.COMPLETED)
        .select_related("customer")
        .prefetch_related(Prefetch("items", queryset=items, to_attr="report_items"))
        .order_by("sale_date", "pk")
    )

    if start:
        sales = sales.filter(sale_date__gte=start)
    if end:
        sales = sales.filter(sale_date__lt=end)
    if product_id:
        sales = sales.filter(items__product_id=product_id).distinct()
    if customer_id:
        sales = sales.filter(customer_id=customer_id)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Completed Sales"
    worksheet.append([
        "Date",
        "Receipt No",
        "Customer",
        "Product",
        "Quantity",
        "Unit Price",
        "Discount",
        "Total",
    ])

    for sale in sales:
        receipt_number = f"SAL-{sale.uuid}"
        discounts = allocate_discount(sale.report_items, sale.discount)
        for item, line_discount in zip(sale.report_items, discounts):
            worksheet.append([
                timezone.localtime(sale.sale_date).date(),
                receipt_number,
                sale.customer.name,
                item.product.name,
                item.quantity,
                item.unit_price,
                line_discount,
                item.subtotal - line_discount,
            ])
            worksheet.cell(worksheet.max_row, 1).number_format = DATE_FORMAT
            worksheet.cell(worksheet.max_row, 5).number_format = "0"
            for column in (6, 7, 8):
                worksheet.cell(worksheet.max_row, column).number_format = MONEY_FORMAT

    style_worksheet(worksheet)
    return excel_response(workbook, report_filename("completed-sales", from_date, to_date))


def get_active_business_details():
    return BusinessDetails.objects.filter(is_active=True).select_related("created_by").first()


def get_payment_summary(sale):
    payments = list(
        sale.payments
        .order_by("payment_date", "pk")
        .values("uuid", "amount", "method", "reference", "payment_date")
    )
    paid_amount = sum((payment["amount"] for payment in payments), Decimal("0.00"))

    return {
        "paid_amount": paid_amount,
        "outstanding_balance": sale.total - paid_amount,
        "payment_methods": sorted({payment["method"] for payment in payments}),
        "payments": [
            {
                "uuid": str(payment["uuid"]),
                "amount": payment["amount"],
                "method": payment["method"],
                "reference": payment["reference"],
                "payment_date": payment["payment_date"].isoformat(),
            }
            for payment in payments
        ],
    }


def build_receipt_data(sale):
    items = list(
        sale.items
        .select_related("product")
        .order_by("created_at", "pk")
        .values(
            "uuid",
            "product__name",
            "quantity",
            "unit_price",
            "subtotal",
        )
    )
    payment_summary = get_payment_summary(sale)
    business = get_active_business_details()
    return {
        "business": {
            "name": business.name if business else "IMARA SHOP",
            "address": business.address if business else "",
            "phone": business.phone if business else "",
            "email": business.email if business else "",
            "tax_number": business.tax_number if business else "",
            "receipt_footer": business.receipt_footer if business else "Thank you for your business.",
        },
        "sale": {
            "uuid": str(sale.uuid),
            "receipt_number": f"SAL-{sale.uuid}",
            "sale_date": timezone.localtime(sale.sale_date).isoformat(),
            "cashier": sale.created_by.full_name or sale.created_by.email,
            "customer": sale.customer.name,
            "payment_status": sale.payment_status,
        },
        "items": [
            {
                "uuid": str(item["uuid"]),
                "product_name": item["product__name"],
                "quantity": item["quantity"],
                "unit_price": item["unit_price"],
                "discount": Decimal("0.00"),
                "line_total": item["subtotal"],
            }
            for item in items
        ],
        "totals": {
            "subtotal": sale.subtotal,
            "discount": sale.discount,
            "grand_total": sale.total,
            "amount_paid": payment_summary["paid_amount"],
            "outstanding_balance": payment_summary["outstanding_balance"],
        },
        "payments": payment_summary["payments"],
        "payment_methods": payment_summary["payment_methods"],
        "currency": "TZS",
    }


def html_text(value):
    return escape(str(value), quote=True)


def money(value):
    return f"{Decimal(value):.2f}"


def receipt_money(value):
    return f"TZS {Decimal(value):,.2f}"


def build_receipt_pdf(receipt):
    """Build a self-contained, printable PDF for a completed sale receipt."""
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=20 * mm,
        leftMargin=20 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title=f"Receipt {receipt['sale']['receipt_number']}",
        author=receipt["business"]["name"],
    )
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name="ReceiptTitle", parent=styles["Heading1"], fontName="Helvetica-Bold",
        fontSize=21, leading=26, textColor=colors.HexColor("#172554"), spaceAfter=3,
    ))
    styles.add(ParagraphStyle(
        name="ReceiptLabel", parent=styles["Normal"], fontName="Helvetica-Bold",
        fontSize=8, leading=10, textColor=colors.HexColor("#64748B"), uppercase=True,
    ))
    styles.add(ParagraphStyle(
        name="ReceiptValue", parent=styles["Normal"], fontSize=10, leading=14,
        textColor=colors.HexColor("#0F172A"),
    ))
    styles.add(ParagraphStyle(
        name="ReceiptRight", parent=styles["Normal"], fontSize=9, leading=12,
        alignment=TA_RIGHT, textColor=colors.HexColor("#334155"),
    ))
    styles.add(ParagraphStyle(
        name="ReceiptCenter", parent=styles["Normal"], fontSize=9, leading=13,
        alignment=TA_CENTER, textColor=colors.HexColor("#64748B"),
    ))

    business = receipt["business"]
    sale = receipt["sale"]
    totals = receipt["totals"]
    sale_date = timezone.localtime(datetime.fromisoformat(sale["sale_date"])).strftime("%d %b %Y, %H:%M")
    payment_methods = ", ".join(method.replace("_", " ").title() for method in receipt["payment_methods"]) or "Not recorded"
    story = [
        Paragraph("SALE RECEIPT", styles["ReceiptLabel"]),
        Paragraph(html_text(business["name"]), styles["ReceiptTitle"]),
    ]
    business_lines = [business["address"], business["phone"], business["email"], business["tax_number"]]
    for line in filter(None, business_lines):
        story.append(Paragraph(html_text(line), styles["ReceiptValue"]))
    story.extend([Spacer(1, 5 * mm), HRFlowable(width="100%", thickness=0.7, color=colors.HexColor("#CBD5E1")), Spacer(1, 4 * mm)])

    metadata = [
        ("Receipt number", sale["receipt_number"]),
        ("Sale date", sale_date),
        ("Customer", sale["customer"]),
        ("Cashier", sale["cashier"] or "-"),
        ("Payment status", sale["payment_status"].title()),
        ("Payment method", payment_methods),
    ]
    metadata_cells = []
    for label, value in metadata:
        metadata_cells.append(Paragraph(html_text(label), styles["ReceiptLabel"]))
        metadata_cells.append(Paragraph(html_text(value), styles["ReceiptValue"]))
    metadata_table = Table([metadata_cells[:4], metadata_cells[4:8], metadata_cells[8:]], colWidths=[27 * mm, 58 * mm, 27 * mm, 58 * mm])
    metadata_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
        ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor("#E2E8F0")),
        ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#E2E8F0")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8),
    ]))
    story.extend([metadata_table, Spacer(1, 7 * mm)])

    item_rows = [["Item", "Qty", "Unit price", "Line total"]]
    for item in receipt["items"]:
        item_rows.append([
            Paragraph(html_text(item["product_name"]), styles["ReceiptValue"]),
            str(item["quantity"]), receipt_money(item["unit_price"]),
            receipt_money(item["line_total"]),
        ])
    item_table = Table(item_rows, colWidths=[85 * mm, 22 * mm, 38 * mm, 45 * mm], repeatRows=1)
    item_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1D4ED8")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 8),
        ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#E2E8F0")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7),
    ]))
    story.extend([item_table, Spacer(1, 6 * mm)])

    total_rows = [["Subtotal", receipt_money(totals["subtotal"])], ["Discount", receipt_money(totals["discount"])], ["Grand total", receipt_money(totals["grand_total"])], ["Amount paid", receipt_money(totals["amount_paid"])], ["Balance", receipt_money(totals["outstanding_balance"])]]
    total_table = Table(total_rows, colWidths=[37 * mm, 42 * mm], hAlign="RIGHT")
    total_table.setStyle(TableStyle([
        ("ALIGN", (1, 0), (1, -1), "RIGHT"), ("FONTNAME", (0, 2), (-1, 2), "Helvetica-Bold"),
        ("FONTSIZE", (0, 2), (-1, 2), 12), ("BACKGROUND", (0, 2), (-1, 2), colors.HexColor("#DBEAFE")),
        ("LINEABOVE", (0, 2), (-1, 2), 0.7, colors.HexColor("#93C5FD")),
        ("LINEBELOW", (0, 2), (-1, 2), 0.7, colors.HexColor("#93C5FD")),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.extend([total_table, Spacer(1, 7 * mm)])

    if receipt["payments"]:
        story.append(Paragraph("PAYMENTS", styles["ReceiptLabel"]))
        payment_rows = [["Method", "Reference", "Date", "Amount"]]
        for payment in receipt["payments"]:
            payment_rows.append([
                payment["method"].replace("_", " ").title(), payment["reference"] or "-",
                timezone.localtime(datetime.fromisoformat(payment["payment_date"])).strftime("%d %b %Y"),
                receipt_money(payment["amount"]),
            ])
        payment_table = Table(payment_rows, colWidths=[40 * mm, 65 * mm, 38 * mm, 40 * mm], repeatRows=1)
        payment_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F1F5F9")), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#E2E8F0")), ("ALIGN", (3, 0), (3, -1), "RIGHT"),
            ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        story.extend([Spacer(1, 3 * mm), payment_table, Spacer(1, 7 * mm)])

    story.extend([HRFlowable(width="100%", thickness=0.7, color=colors.HexColor("#CBD5E1")), Spacer(1, 4 * mm), Paragraph(html_text(business["receipt_footer"] or "Thank you for your business."), styles["ReceiptCenter"])])
    document.build(story)
    output.seek(0)
    return output


def build_receipt_html(receipt):
    business = receipt["business"]
    sale = receipt["sale"]
    totals = receipt["totals"]
    payment_methods = ", ".join(receipt["payment_methods"]) or "Not recorded"
    payments = receipt["payments"]

    business_lines = [
        business["name"],
        business["address"],
        business["phone"],
        business["email"],
        business["tax_number"],
    ]
    item_rows = "".join(
        "<tr>"
        f"<td>{html_text(item['product_name'])}</td>"
        f"<td class=\"numeric\">{html_text(item['quantity'])}</td>"
        f"<td class=\"numeric\">{money(item['unit_price'])}</td>"
        f"<td class=\"numeric\">{money(item['line_total'])}</td>"
        "</tr>"
        for item in receipt["items"]
    )
    payment_rows = "".join(
        "<tr>"
        f"<td>{html_text(str(item['method']).replace('_', ' ').title())}</td>"
        f"<td>{html_text(item['reference'] or '—')}</td>"
        f"<td class=\"numeric\">{money(item['amount'])}</td>"
        "</tr>"
        for item in payments
    )
    business_html = (
        f"<strong>{html_text(business['name'])}</strong>"
        + "".join(
            f"<span>{html_text(line)}</span>"
            for line in business_lines[1:]
            if line
        )
    )
    sale_date = datetime.fromisoformat(sale["sale_date"])
    sale_date = timezone.localtime(sale_date).strftime("%d/%m/%Y %H:%M")
    payments_html = ""
    if payment_rows:
        payments_html = (
            "<h3>Payments</h3>"
            "<table><thead><tr><th>Method</th><th>Reference</th>"
            "<th class=\"numeric\">Amount</th></tr></thead><tbody>"
            f"{payment_rows}</tbody></table>"
        )

    return f"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Receipt {html_text(sale['receipt_number'])}</title>
<style>
:root {{ color-scheme: light; --ink: #111827; --muted: #6b7280; --line: #d1d5db; }}
* {{ box-sizing: border-box; }}
body {{ margin: 0; background: #e5e7eb; color: var(--ink); font: 13px/1.45 Arial, sans-serif; }}
.receipt {{ width: 80mm; min-height: 100vh; margin: 0 auto; padding: 18px 14px; background: #fff; }}
.center {{ text-align: center; }}
.business {{ margin-bottom: 16px; }}
.business strong {{ display: block; font-size: 18px; letter-spacing: .02em; }}
.business span {{ display: block; color: var(--muted); }}
.meta {{ display: grid; grid-template-columns: 1fr 1fr; gap: 8px; margin: 14px 0; padding: 10px 0; border-top: 1px solid var(--line); border-bottom: 1px solid var(--line); }}
.meta div:nth-child(odd) {{ color: var(--muted); }}
table {{ width: 100%; border-collapse: collapse; font-variant-numeric: tabular-nums; }}
th {{ padding: 6px 3px; color: var(--muted); font-size: 11px; text-align: left; border-bottom: 1px solid var(--line); }}
td {{ padding: 7px 3px; border-bottom: 1px solid #e5e7eb; }}
.numeric {{ text-align: right; white-space: nowrap; }}
.totals {{ width: 65%; margin-left: auto; }}
.totals td {{ padding: 4px 3px; }}
.totals .grand {{ border-top: 1px solid var(--line); border-bottom: 1px solid var(--line); font-size: 16px; font-weight: 700; }}
.payments {{ margin-top: 14px; font-size: 12px; }}
.footer {{ margin-top: 20px; padding-top: 12px; border-top: 1px solid var(--line); }}
.actions {{ position: fixed; right: 16px; bottom: 16px; display: flex; gap: 8px; }}
button {{ border: 0; border-radius: 6px; padding: 10px 14px; background: #1d4ed8; color: #fff; font: 600 13px Arial, sans-serif; cursor: pointer; }}
@media print {{
  @page {{ size: 80mm auto; margin: 8mm; }}
  body {{ background: #fff; font-size: 12px; }}
  .receipt {{ width: 80mm; min-height: 0; margin: 0; padding: 0; }}
  .actions {{ display: none; }}
  a {{ color: inherit; text-decoration: none; }}
}}
@media print and (min-width: 1000px) {{
  @page {{ size: A4; margin: 16mm; }}
  .receipt {{ width: 130mm; padding: 12px; }}
}}
</style>
</head>
<body>
<main class="receipt">
  <section class="business center">{business_html}</section>
  <div class="meta">
    <div>Receipt</div><div><strong>{html_text(sale['receipt_number'])}</strong></div>
    <div>Date</div><div>{sale_date}</div>
    <div>Cashier</div><div>{html_text(sale['cashier'])}</div>
    <div>Customer</div><div>{html_text(sale['customer'])}</div>
    <div>Payment status</div><div>{html_text(sale['payment_status'].title())}</div>
    <div>Payment method</div><div>{html_text(payment_methods)}</div>
  </div>
  <table>
    <thead><tr><th>Item</th><th class="numeric">Qty</th><th class="numeric">Price</th><th class="numeric">Total</th></tr></thead>
    <tbody>{item_rows}</tbody>
  </table>
  <table class="totals">
    <tr><td>Subtotal</td><td class="numeric">{money(totals['subtotal'])}</td></tr>
    <tr><td>Discount</td><td class="numeric">{money(totals['discount'])}</td></tr>
    <tr class="grand"><td>Total</td><td class="numeric">{money(totals['grand_total'])}</td></tr>
    <tr><td>Paid</td><td class="numeric">{money(totals['amount_paid'])}</td></tr>
    <tr><td>Balance</td><td class="numeric">{money(totals['outstanding_balance'])}</td></tr>
  </table>
  {payments_html}
  <footer class="footer center">{html_text(business['receipt_footer'])}</footer>
</main>
<div class="actions"><button type="button" onclick="window.print()">Print receipt</button></div>
</body>
</html>"""
